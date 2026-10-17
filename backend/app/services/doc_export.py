"""Export the grant editor document (TipTap HTML) to PDF or DOCX.

PDF renders the HTML in headless Chromium (Playwright, already used by the
scrapers) so colours/fonts/tables match what the user sees. DOCX walks the HTML
with BeautifulSoup and rebuilds it with python-docx, carrying headings, inline
bold/italic/underline, text colour, font family/size, lists, and tables.
"""
from __future__ import annotations

import io
import re

from bs4 import BeautifulSoup, NavigableString, Tag

# Wraps the document body with print CSS approximating the editor's paper.
_PDF_TEMPLATE = """<!DOCTYPE html><html><head><meta charset="utf-8">
<style>
  @page {{ margin: 2.5cm 2cm; }}
  body {{ font-family: Georgia, 'Times New Roman', serif; font-size: 11pt; color: #1f2937; line-height: 1.5; }}
  h1 {{ font-size: 20pt; margin: 0 0 10px; }}
  h2 {{ font-size: 15pt; margin: 18px 0 6px; border-bottom: 1px solid #e5e7eb; padding-bottom: 2px; }}
  h3 {{ font-size: 13pt; margin: 14px 0 4px; }}
  h4 {{ font-size: 12pt; margin: 12px 0 4px; }}
  p {{ margin: 6px 0; }}
  table {{ border-collapse: collapse; width: 100%; }}
  td, th {{ border: 1px solid #d1d5db; padding: 4px 6px; }}
  th {{ background: #f9fafb; }}
  img {{ max-width: 100%; height: auto; }}
  a {{ color: #2563eb; }}
  /* Nested list markers by depth — mirrors the editor (globals.css). */
  ul {{ list-style-type: disc; }}
  :is(ul, ol) ul {{ list-style-type: circle; }}
  :is(ul, ol) :is(ul, ol) ul {{ list-style-type: square; }}
  ol {{ list-style-type: decimal; }}
  :is(ul, ol) ol {{ list-style-type: lower-alpha; }}
  :is(ul, ol) :is(ul, ol) ol {{ list-style-type: lower-roman; }}
  li > p {{ margin: 2px 0; }}
  div[data-type="page-break"] {{ page-break-after: always; border: 0; height: 0; visibility: hidden; }}
</style></head><body>{body}</body></html>"""


def html_to_pdf(content_html: str) -> bytes:
    """Render HTML → PDF bytes via headless Chromium. Runs the sync Playwright
    API, so call it from a thread (asyncio.to_thread) inside async code."""
    from playwright.sync_api import sync_playwright

    full = _PDF_TEMPLATE.format(body=content_html or "")
    with sync_playwright() as p:
        browser = p.chromium.launch(args=["--no-sandbox"])
        try:
            page = browser.new_page()
            page.set_content(full, wait_until="load")
            pdf = page.pdf(format="A4", print_background=True,
                           margin={"top": "0", "bottom": "0", "left": "0", "right": "0"})
        finally:
            browser.close()
    return pdf


# ── DOCX ────────────────────────────────────────────────────────────────────

def _apply_style_attrs(run, style: str) -> None:
    """Apply inline CSS (color, font-family, font-size) from a span to a docx run."""
    from docx.shared import Pt, RGBColor

    if not style:
        return
    decls = dict(
        (k.strip().lower(), v.strip())
        for k, v in (part.split(":", 1) for part in style.split(";") if ":" in part)
    )
    color = decls.get("color")
    if color and color.startswith("#") and len(color) in (4, 7):
        hexv = color[1:]
        if len(hexv) == 3:
            hexv = "".join(c * 2 for c in hexv)
        try:
            run.font.color.rgb = RGBColor.from_string(hexv.upper())
        except ValueError:
            pass
    fam = decls.get("font-family")
    if fam:
        run.font.name = fam.split(",")[0].strip().strip('"\'')
    size = decls.get("font-size")
    if size:
        m = re.match(r"([\d.]+)\s*(pt|px)?", size)
        if m:
            val = float(m.group(1))
            if (m.group(2) or "pt") == "px":
                val = val * 0.75  # px → pt
            run.font.size = Pt(val)


def _add_runs(paragraph, node, *, bold=False, italic=False, underline=False, style="") -> None:
    """Recursively emit docx runs from an inline HTML subtree."""
    for child in node.children:
        if isinstance(child, NavigableString):
            text = str(child)
            if text:
                run = paragraph.add_run(text)
                run.bold = bold or None
                run.italic = italic or None
                run.underline = underline or None
                _apply_style_attrs(run, style)
        elif isinstance(child, Tag):
            b = bold or child.name in ("b", "strong")
            i = italic or child.name in ("i", "em")
            u = underline or child.name in ("u",)
            s = child.get("style", "") or style
            if child.name == "br":
                paragraph.add_run().add_break()
            else:
                _add_runs(paragraph, child, bold=b, italic=i, underline=u, style=s)


_HEADING_STYLE = {"h1": "Title", "h2": "Heading 1", "h3": "Heading 2", "h4": "Heading 3"}


def _add_table(doc, table_tag) -> None:
    rows = table_tag.find_all("tr")
    if not rows:
        return
    ncols = max(len(r.find_all(["td", "th"])) for r in rows)
    table = doc.add_table(rows=0, cols=ncols)
    table.style = "Table Grid"
    for r in rows:
        cells = r.find_all(["td", "th"])
        row = table.add_row().cells
        for idx, cell in enumerate(cells):
            if idx < ncols:
                row[idx].text = cell.get_text(" ", strip=True)


# Per-level list markers, matching the editor and the PDF export. The default
# python-docx template's "List Bullet 2/3" styles all use the same glyph, so we
# register our own multi-level numbering definitions instead.
_BULLET_GLYPHS = ["\u2022", "\u25e6", "\u25aa"]  # • ◦ ▪
_NUMBER_FORMATS = ["decimal", "lowerLetter", "lowerRoman"]
_LIST_LEVELS = 9  # Word supports ilvl 0-8


class _ListNumbering:
    """Creates the bullet / ordered abstract numbering once per document and hands
    out numIds: one shared id for bullets, a fresh restarting id per <ol>."""

    def __init__(self, doc):
        from docx.oxml import parse_xml
        from docx.oxml.ns import nsdecls

        self._numbering = doc.part.numbering_part.element
        existing = [int(a) for a in self._numbering.xpath("./w:abstractNum/@w:abstractNumId")]
        next_id = max(existing, default=0) + 1
        self._abstract = {}
        for kind in ("ul", "ol"):
            levels = []
            for i in range(_LIST_LEVELS):
                if kind == "ul":
                    fmt, text = "bullet", _BULLET_GLYPHS[i % len(_BULLET_GLYPHS)]
                else:
                    fmt, text = _NUMBER_FORMATS[i % len(_NUMBER_FORMATS)], f"%{i + 1}."
                levels.append(
                    f'<w:lvl w:ilvl="{i}"><w:start w:val="1"/><w:numFmt w:val="{fmt}"/>'
                    f'<w:lvlText w:val="{text}"/><w:lvlJc w:val="left"/>'
                    f'<w:pPr><w:ind w:left="{360 * (i + 2)}" w:hanging="360"/></w:pPr></w:lvl>'
                )
            el = parse_xml(
                f'<w:abstractNum {nsdecls("w")} w:abstractNumId="{next_id}">'
                f'<w:multiLevelType w:val="hybridMultilevel"/>{"".join(levels)}</w:abstractNum>'
            )
            # abstractNum elements must precede the num elements.
            first_num = self._numbering.find("{%s}num" % el.nsmap["w"])
            if first_num is not None:
                first_num.addprevious(el)
            else:
                self._numbering.append(el)
            self._abstract[kind] = next_id
            next_id += 1
        self._bullet_num = self._numbering.add_num(self._abstract["ul"]).numId

    def num_id(self, list_tag: str, level: int) -> int:
        if list_tag == "ul":
            return self._bullet_num
        num = self._numbering.add_num(self._abstract["ol"])
        num.add_lvlOverride(ilvl=level).add_startOverride(1)
        return num.numId


def _add_list(doc, numbering: _ListNumbering, list_el, level: int = 0) -> None:
    """Emit one paragraph per <li>, recursing into nested <ul>/<ol> one level deeper
    so indentation and per-level markers survive the export."""
    num_id = numbering.num_id(list_el.name, min(level, _LIST_LEVELS - 1))
    for li in list_el.find_all("li", recursive=False):
        # Detach nested lists so the item's own text doesn't swallow them.
        nested = [c.extract() for c in li.find_all(["ul", "ol"], recursive=False)]
        p = doc.add_paragraph(style="List Paragraph")
        num_pr = p._p.get_or_add_pPr().get_or_add_numPr()
        num_pr.get_or_add_ilvl().val = min(level, _LIST_LEVELS - 1)
        num_pr.get_or_add_numId().val = num_id
        paras = li.find_all("p", recursive=False)
        if paras:
            # TipTap wraps item text in <p>; extra paragraphs share the bullet.
            for i, para in enumerate(paras):
                if i:
                    p.add_run().add_break()
                _add_runs(p, para, style=para.get("style", "") or li.get("style", ""))
        else:
            _add_runs(p, li, style=li.get("style", ""))
        for sub in nested:
            _add_list(doc, numbering, sub, level + 1)


def html_to_docx(content_html: str) -> bytes:
    """Convert TipTap HTML → DOCX bytes."""
    from docx import Document

    soup = BeautifulSoup(content_html or "", "html.parser")
    doc = Document()
    numbering = _ListNumbering(doc)
    root = soup.body or soup

    def walk(container):
        for el in container.children:
            if isinstance(el, NavigableString):
                text = str(el).strip()
                if text:
                    doc.add_paragraph().add_run(text)
                continue
            if not isinstance(el, Tag):
                continue
            name = el.name
            if name in _HEADING_STYLE:
                p = doc.add_paragraph(style=_HEADING_STYLE[name])
                _add_runs(p, el, style=el.get("style", ""))
            elif name == "p":
                p = doc.add_paragraph()
                _add_runs(p, el, style=el.get("style", ""))
            elif name in ("ul", "ol"):
                _add_list(doc, numbering, el)
            elif name == "blockquote":
                p = doc.add_paragraph(style="Intense Quote")
                _add_runs(p, el, style=el.get("style", ""))
            elif name == "table":
                _add_table(doc, el)
            elif name in ("div", "section", "article"):
                if el.get("data-type") == "page-break":
                    doc.add_page_break()
                else:
                    walk(el)  # descend into wrappers
            else:
                p = doc.add_paragraph()
                _add_runs(p, el, style=el.get("style", ""))

    walk(root)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
